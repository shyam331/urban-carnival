"""Write campaign_zips.xlsx.

All ZIPs carries live formulas for tier, weight and scores that read the thresholds on
the Assumptions tab, so editing a tier floor there re-tiers every ZIP in Excel. Summary
and Valpak are sorted snapshots produced by the script at the configured defaults
(Excel cannot re-sort without dynamic-array functions); re-run the script after
changing config.py to refresh them.
"""
import pandas as pd
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from . import config as C

FONT = "Arial"
HDR_FILL = PatternFill("solid", fgColor="1F3864")
INPUT_FILL = PatternFill("solid", fgColor="FFFF00")
NOTE_FONT = Font(name=FONT, italic=True, size=9, color="595959")
WARN_FONT = Font(name=FONT, bold=True, color="C00000")
NO_DATA = "no data"

# (header, source column or None for formula, number format)
ALL_COLS = [
    ("Rank", "rank", "0"),
    ("ZIP", "zip", "@"),
    ("City", "city", None),
    ("County", "county", None),
    ("Focus region", "focus_region", None),
    ("Neighborhoods", "neighborhoods", None),
    ("Redfin metro", "metro", None),
    ("Median sale $/sq ft (12 mo)", "ppsf_12m", "$#,##0"),
    ("$/sq ft basis", "ppsf_basis", None),
    ("Median sale price (12 mo)", "sale_price_12m", "$#,##0"),
    ("Homes sold (12 mo)", "homes_sold_12m", "#,##0"),
    ("Thin sample (<%d sales)" % C.MIN_HOMES_SOLD_12M, "thin_sample", None),
    ("Has ZIP boundary (ZCTA)", "has_zcta", None),
    ("Tier", None, "0"),
    ("Tier weight", None, "0"),
    ("Total SFR parcels", "sfr_parcels", "#,##0"),
    ("SFR lots >= %s sq ft" % f"{C.MIN_LOT_SQFT:,}", "sfr_lot_5k", "#,##0"),
    (f"% SFR lots >= {C.MIN_LOT_SQFT:,}", "pct_lot_5k", "0.0%"),
    ("Est. buildable >= %d sq ft" % C.ADU_MIN_BUILDABLE_SQFT, "buildable_750", "#,##0"),
    (f"% est. buildable >= {C.ADU_MIN_BUILDABLE_SQFT}", "pct_buildable_750", "0.0%"),
    ("Est. buildable >= %s sq ft" % f"{C.ADU_LARGE_BUILDABLE_SQFT:,}", "buildable_1200", "#,##0"),
    ("Flag: contiguous area likely < %d" % C.ADU_MIN_BUILDABLE_SQFT, "contiguous_under_750_flag", "#,##0"),
    ("Median lot size (sq ft)", "median_lot_sqft", "#,##0"),
    ("Parcel data", "parcel_data", None),
    ("Rank score (qualifying parcels x weight)", None, "#,##0"),
    ("Interim score (homes sold x weight)", None, "#,##0"),
    ("Rank basis", "rank_basis", None),
    ("Building footprints (all)", "footprints_total", "#,##0"),
    ("House-scale footprints (600-6,000 sq ft)", "footprints_house_scale", "#,##0"),
    ("Footprints sourced from Microsoft", "footprints_microsoft", "#,##0"),
    ("Median house-scale footprint (sq ft)", "median_house_scale_footprint_sqft", "#,##0"),
    ("Latest-quarter $/sq ft", "sfr_latest_qtr_ppsf", "$#,##0"),
    ("Latest-quarter $/sq ft YoY", "sfr_latest_qtr_ppsf_yoy", "0.0%"),
    ("County share of ZIP area", "county_share", "0%"),
    ("City source", "city_source", None),
    ("Other postal cities (>=10% of addresses)", "postal_cities", None),
    ("Overture locality (largest overlap)", "locality", None),
]


def _hdr(ws, row, headers, widths=None):
    for j, h in enumerate(headers, 1):
        c = ws.cell(row=row, column=j, value=h)
        c.font = Font(name=FONT, bold=True, color="FFFFFF")
        c.fill = HDR_FILL
        c.alignment = Alignment(wrap_text=True, vertical="center")
        if widths:
            ws.column_dimensions[get_column_letter(j)].width = widths[j - 1] if isinstance(widths, list) else widths
    ws.row_dimensions[row].height = 42


def _val(v):
    if v is None or (not isinstance(v, str) and pd.isna(v)):
        return None
    if hasattr(v, "item"):
        return v.item()
    return v


def _font_all(ws):
    for row in ws.iter_rows():
        for c in row:
            if c.font is None or c.font.name != FONT:
                f = c.font
                c.font = Font(name=FONT, bold=f.bold, italic=f.italic, color=f.color, size=f.size or 10)


def write_assumptions(wb, meta):
    ws = wb.create_sheet("Assumptions")
    ws["A1"] = "Assumptions and configurable thresholds"
    ws["A1"].font = Font(name=FONT, bold=True, size=14)
    ws["A2"] = ("Yellow cells drive the live formulas on 'All ZIPs' (tier, weight, scores) and 'By County'. "
                "Every value also lives in adu_campaign/config.py; re-run the script to refresh Summary and Valpak.")
    ws["A2"].font = NOTE_FONT
    _hdr(ws, 4, ["Parameter", "Value", "Notes / source"], [44, 16, 110])
    rows = [
        ("Tier 1 floor ($/sq ft)", C.PRICE_TIERS[0]["min_ppsf"], "Brief, criterion 4. Redfin median sale $/sq ft, trailing 12 months"),
        ("Tier 2 floor ($/sq ft)", C.PRICE_TIERS[1]["min_ppsf"], "Brief: second tier"),
        ("Tier 3 floor ($/sq ft)", C.PRICE_TIERS[2]["min_ppsf"], "Brief: third tier"),
        ("Tier 1 weight", C.PRICE_TIERS[0]["weight"], "Rank score multiplier (our choice: 3/2/1)"),
        ("Tier 2 weight", C.PRICE_TIERS[1]["weight"], "Rank score multiplier"),
        ("Tier 3 weight", C.PRICE_TIERS[2]["weight"], "Rank score multiplier"),
        ("Minimum lot size (sq ft)", C.MIN_LOT_SQFT, "Brief, criterion 2"),
        ("Side setback (ft)", C.SIDE_SETBACK_FT, "CA state ADU baseline, Gov. Code 66314(d)(7) (formerly 65852.2)"),
        ("Rear setback (ft)", C.REAR_SETBACK_FT, "CA state ADU baseline"),
        ("Front setback (ft)", C.FRONT_SETBACK_FT, "Default assumption from the brief; local codes vary (often 15-25 ft)"),
        ("Separation from existing house (ft)", C.BUILDING_SEPARATION_FT, "0 = none assumed (attached ADUs allowed by state law)"),
        ("ADU buildable threshold (sq ft)", C.ADU_MIN_BUILDABLE_SQFT, "Brief, criterion 3"),
        ("Large ADU band (sq ft)", C.ADU_LARGE_BUILDABLE_SQFT, "Brief: 'even 1,200 sq ft or more' reported as a separate count"),
        ("Footprint-in-parcel overlap minimum", C.FOOTPRINT_PARCEL_OVERLAP_MIN, "Share of a footprint's area that must fall inside a parcel to count as that parcel's building"),
        ("Thin-sample flag (homes sold, 12 mo)", C.MIN_HOMES_SOLD_12M, "ZIPs below this are flagged; tier still assigned"),
        ("Redfin property type", C.REDFIN_PROPERTY_TYPE, f"Fallback: {C.REDFIN_FALLBACK_PROPERTY_TYPE} (flagged in '$/sq ft basis')"),
        ("Redfin windows used", ", ".join(meta["redfin"]["windows"]), "Four non-overlapping 90-day windows; $/sf and price are homes-sold-weighted averages of window medians"),
        ("Redfin file last updated", str(meta["redfin"]["last_updated"]), "Latest period in the public file; it was ~5 months old at build time"),
        ("Overture release", C.OVERTURE_RELEASE, "Building footprints (incl. Microsoft ML) and county/city/neighborhood names"),
        ("ZIP boundaries", meta["zcta"]["url"], "2010 TIGER ZCTA mirror (Census host unreachable from the build environment)"),
        ("Interim ranking", "Redfin SFR homes sold (12 mo) x tier weight", "Used ONLY where parcel data is missing; replace by re-running with parcel data"),
    ]
    for i, (k, v, n) in enumerate(rows, 5):
        ws.cell(row=i, column=1, value=k)
        c = ws.cell(row=i, column=2, value=v)
        if isinstance(v, (int, float)):
            c.fill = INPUT_FILL
            c.font = Font(name=FONT, color="0000FF")
        ws.cell(row=i, column=3, value=n)
    return {"t1": "Assumptions!$B$5", "t2": "Assumptions!$B$6", "t3": "Assumptions!$B$7",
            "w1": "Assumptions!$B$8", "w2": "Assumptions!$B$9", "w3": "Assumptions!$B$10"}


def write_all_zips(wb, z, refs):
    ws = wb.create_sheet("All ZIPs")
    heads = [h for h, _, _ in ALL_COLS]
    _hdr(ws, 1, heads, 14)
    ws.column_dimensions["F"].width = 40
    col = {h: get_column_letter(j) for j, h in enumerate(heads, 1)}
    P, T, W = col["Median sale $/sq ft (12 mo)"], col["Tier"], col["Tier weight"]
    B750, PD, HS = col["Est. buildable >= %d sq ft" % C.ADU_MIN_BUILDABLE_SQFT], col["Parcel data"], col["Homes sold (12 mo)"]
    for i, r in enumerate(z.itertuples(index=False), 2):
        rd = r._asdict()
        for j, (h, src, fmt) in enumerate(ALL_COLS, 1):
            if src is not None:
                v = _val(rd.get(src))
                if v is None and src in ("sfr_parcels", "sfr_lot_5k", "pct_lot_5k", "buildable_750", "pct_buildable_750",
                                         "buildable_1200", "contiguous_under_750_flag", "median_lot_sqft"):
                    v = NO_DATA
            elif h == "Tier":
                v = (f'=IF({P}{i}="","",IF({P}{i}>={refs["t1"]},1,IF({P}{i}>={refs["t2"]},2,'
                     f'IF({P}{i}>={refs["t3"]},3,""))))')
            elif h == "Tier weight":
                v = f'=IF({T}{i}=1,{refs["w1"]},IF({T}{i}=2,{refs["w2"]},IF({T}{i}=3,{refs["w3"]},0)))'
            elif h.startswith("Rank score"):
                v = f'=IF(AND({PD}{i}="processed",ISNUMBER({B750}{i})),{B750}{i}*{W}{i},"")'
            elif h.startswith("Interim"):
                v = f'=IF(ISNUMBER({HS}{i}),{HS}{i}*{W}{i},0)'
            c = ws.cell(row=i, column=j, value=v)
            if fmt and fmt != "@":
                c.number_format = fmt
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(heads))}{len(z) + 1}"
    return col, len(z) + 1


def write_summary(wb, z, meta, col, last):
    ws = wb["Summary"]
    ws["A1"] = "California ADU direct-mail targeting: top %d ZIPs" % C.SUMMARY_TOP_N
    ws["A1"].font = Font(name=FONT, bold=True, size=14)
    any_parcels = (z.parcel_data == "processed").any()
    ws["A2"] = ("Ranking = qualifying parcels x tier weight." if any_parcels else
                "INTERIM RANKING: no county parcel data could be downloaded in this run (assessor hosts unreachable). "
                "Rank = Redfin SFR homes sold (12 mo) x tier weight, a market-size proxy, NOT a count of qualifying lots. "
                "Parcel columns read 'no data'.")
    ws["A2"].font = WARN_FONT
    ws["A3"] = ("ZIP-level results are estimates of the share of qualifying lots, not a verified list of homes. "
                f"Redfin data: {', '.join(meta['redfin']['windows'])} windows (file updated {meta['redfin']['last_updated']}).")
    ws["A3"].font = NOTE_FONT

    heads5 = ["Tier", "Mailable ZIPs (live count)", "Cumulative ZIPs", "SFR homes sold 12 mo in tier",
              "PO-box / unique ZIPs in tier (excluded)"]
    for j, h in enumerate(heads5, 1):
        ws.cell(row=5, column=j, value=h).font = Font(name=FONT, bold=True)
    tcol, hcol, mcol = col["Tier"], col["Homes sold (12 mo)"], col["Has ZIP boundary (ZCTA)"]
    rng = lambda c: f"'All ZIPs'!${c}$2:${c}${last}"  # noqa: E731
    for k, t in enumerate(C.PRICE_TIERS, 6):
        ws.cell(row=k, column=1, value=t["label"])
        ws.cell(row=k, column=2, value=f"=COUNTIFS({rng(tcol)},{t['tier']},{rng(mcol)},TRUE)")
        ws.cell(row=k, column=3, value=f"=SUM($B$6:B{k})")
        ws.cell(row=k, column=4, value=f"=SUMIFS({rng(hcol)},{rng(tcol)},{t['tier']},{rng(mcol)},TRUE)").number_format = "#,##0"
        ws.cell(row=k, column=5, value=f"=COUNTIFS({rng(tcol)},{t['tier']},{rng(mcol)},FALSE)")
    r = 6 + len(C.PRICE_TIERS)
    ws.cell(row=r, column=1, value="Total (any tier)").font = Font(name=FONT, bold=True)
    ws.cell(row=r, column=2, value=f"=SUM(B6:B{r - 1})")
    ws.cell(row=r, column=4, value=f"=SUM(D6:D{r - 1})").number_format = "#,##0"
    ws.cell(row=r, column=5, value=f"=SUM(E6:E{r - 1})")

    start = r + 2
    heads = ["Rank", "ZIP", "City", "County", "Focus region", "Neighborhoods", "Tier", "Median $/sq ft (12 mo)",
             "Median sale price (12 mo)", "SFR homes sold (12 mo)", "Est. buildable >= 750 (parcels)", "Score", "Rank basis"]
    _hdr(ws, start, heads, [7, 8, 18, 15, 26, 45, 18, 12, 14, 12, 14, 10, 34])
    top = z[z["rank"].notna()].head(C.SUMMARY_TOP_N)
    for i, rr in enumerate(top.itertuples(index=False), start + 1):
        vals = [rr.rank, rr.zip, rr.city, rr.county, rr.focus_region, rr.neighborhoods, rr.tier_label, rr.ppsf_12m,
                rr.sale_price_12m, rr.homes_sold_12m, NO_DATA if rr.parcel_data != "processed" else rr.buildable_750,
                rr.score, rr.rank_basis]
        for j, v in enumerate(vals, 1):
            c = ws.cell(row=i, column=j, value=_val(v))
            c.number_format = {8: "$#,##0", 9: "$#,##0", 10: "#,##0", 12: "#,##0"}.get(j, "General")
    ws.freeze_panes = ws.cell(row=start + 1, column=3)
    ws.cell(row=start - 1, column=1, value="Snapshot sorted by the script at config.py defaults.").font = NOTE_FONT


def write_by_county(wb, z, col, last, inv):
    ws = wb.create_sheet("By County")
    ws["A1"] = "Focus regions and every California county (live counts from 'All ZIPs')"
    ws["A1"].font = Font(name=FONT, bold=True, size=13)
    heads = ["Region / county", "Counties", "ZIPs", "Tier 1 ZIPs", "Tier 2 ZIPs", "Tier 3 ZIPs", "Any-tier ZIPs",
             "SFR homes sold 12 mo (tiered ZIPs)", "Median of ZIP $/sq ft", "Parcel data status"]
    _hdr(ws, 3, heads, [44, 30, 8, 10, 10, 10, 12, 18, 14, 60])
    T, R, Cc, H, P = col["Tier"], col["Focus region"], col["County"], col["Homes sold (12 mo)"], col["Median sale $/sq ft (12 mo)"]
    rng = lambda c: f"'All ZIPs'!${c}$2:${c}${last}"  # noqa: E731
    status = dict(zip(inv.county, inv.status_from_build_env))

    def row(i, label, counties, key_col, key):
        ws.cell(row=i, column=1, value=label)
        ws.cell(row=i, column=2, value=", ".join(counties))
        ws.cell(row=i, column=3, value=f'=COUNTIF({rng(key_col)},"{key}")')
        for k, t in enumerate((1, 2, 3), 4):
            ws.cell(row=i, column=k, value=f'=COUNTIFS({rng(key_col)},"{key}",{rng(T)},{t})')
        ws.cell(row=i, column=7, value=f"=SUM(D{i}:F{i})")
        ws.cell(row=i, column=8, value=f'=SUMIFS({rng(H)},{rng(key_col)},"{key}",{rng(T)},">0")').number_format = "#,##0"
        med = z[z[{R: "focus_region", Cc: "county"}[key_col]] == key].ppsf_12m.median()
        ws.cell(row=i, column=9, value=_val(med)).number_format = "$#,##0"
        ws.cell(row=i, column=10, value="; ".join(f"{c}: {status.get(c, 'n/a')}" for c in counties)[:500])

    i = 4
    ws.cell(row=i, column=1, value="FOCUS REGIONS").font = Font(name=FONT, bold=True)
    i += 1
    for region, counties in C.FOCUS_REGIONS.items():
        row(i, region, counties, R, region)
        i += 1
    row(i, "Rest of California", ["all other counties"], R, "Rest of California")
    i += 2
    ws.cell(row=i, column=1, value="ALL COUNTIES").font = Font(name=FONT, bold=True)
    i += 1
    for county in sorted(z.county.dropna().unique()):
        row(i, county, [county], Cc, county)
        i += 1
    ws.cell(row=i + 1, column=1, value="Median of ZIP $/sq ft is a static value at build time; other counts are live formulas.").font = NOTE_FONT
    ws.freeze_panes = "B4"


def write_neighborhoods(wb, z, hoods):
    ws = wb.create_sheet("Neighborhoods")
    h = hoods.merge(z[["zip", "city", "county", "tier_label", "rank"]], on="zip", how="left")
    h = h.sort_values(["rank", "zip", "neighborhood"], na_position="last")
    heads = ["ZIP", "City", "County", "Neighborhood", "Overture place type", "Tier", "ZIP rank"]
    _hdr(ws, 1, heads, [8, 20, 16, 32, 16, 20, 9])
    for i, r in enumerate(h.itertuples(index=False), 2):
        for j, v in enumerate([r.zip, r.city, r.county, r.neighborhood, r.subtype, r.tier_label, r.rank], 1):
            ws.cell(row=i, column=j, value=_val(v))
    n = len(h) + 2
    ws.cell(row=n + 1, column=1, value=("Neighborhood points (Overture divisions: macrohood / neighborhood / microhood) "
                                        "joined to ZIP polygons by point-in-polygon. ZIPs with no named neighborhood "
                                        "are listed on 'All ZIPs' with a blank Neighborhoods cell.")).font = NOTE_FONT
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = f"A1:G{len(h) + 1}"


def write_valpak(wb, z):
    ws = wb.create_sheet("Valpak")
    heads = ["ZIP", "City", "Neighborhood", "Qualifying household estimate", "Tier",
             "Rank", "SFR homes sold (12 mo)", "Median $/sq ft (12 mo)", "Estimate basis"]
    _hdr(ws, 1, heads, [8, 20, 40, 18, 20, 7, 12, 12, 60])
    v = z[z["rank"].notna()]
    for i, r in enumerate(v.itertuples(index=False), 2):
        hoods = "; ".join(str(r.neighborhoods).split("; ")[:3]) if isinstance(r.neighborhoods, str) else None
        if r.parcel_data == "processed":
            est, basis = _val(r.buildable_750), "SFR parcels, lot >= 5,000 sq ft, est. buildable >= 750 sq ft"
        else:
            est, basis = NO_DATA, "Parcel data unavailable for this county; count pending (see Source Inventory)"
        for j, val in enumerate([r.zip, r.city, hoods, est, r.tier_label, r.rank, r.homes_sold_12m, r.ppsf_12m, basis], 1):
            c = ws.cell(row=i, column=j, value=_val(val))
            c.number_format = {7: "#,##0", 8: "$#,##0", 4: "#,##0"}.get(j, "General")
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = f"A1:I{len(v) + 1}"


def write_inventory(wb, inv):
    ws = wb.create_sheet("Source Inventory")
    heads = ["County", "Source", "URL", "Status from build environment", "Data vintage", "Lot size field",
             "Use code field", "SFR rule", "Coverage %", "Notes"]
    _hdr(ws, 1, heads, [16, 44, 60, 40, 30, 26, 18, 16, 22, 80])
    for i, r in enumerate(inv.itertuples(index=False), 2):
        for j, v in enumerate(r, 1):
            ws.cell(row=i, column=j, value=_val(v))
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = f"A1:J{len(inv) + 1}"
    ws.cell(row=1, column=4).comment = Comment("Probed by adu_campaign.util.probe at build time. "
                                               "UNREACHABLE means the host was blocked or did not answer from the build "
                                               "environment; it does not mean the source is offline.", "pipeline")


def write_workbook(path, z, hoods, inv, meta):
    wb = Workbook()
    wb.active.title = "Summary"
    refs = write_assumptions(wb, meta)
    col, last = write_all_zips(wb, z, refs)
    write_summary(wb, z, meta, col, last)
    write_by_county(wb, z, col, last, inv)
    write_neighborhoods(wb, z, hoods)
    write_valpak(wb, z)
    write_inventory(wb, inv)
    order = ["Summary", "All ZIPs", "By County", "Neighborhoods", "Valpak", "Source Inventory", "Assumptions"]
    wb._sheets = [wb[n] for n in order]
    for ws in wb.worksheets:
        _font_all(ws)
    wb.save(path)
